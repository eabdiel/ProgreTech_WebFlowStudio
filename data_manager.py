from __future__ import annotations

from pathlib import Path
from threading import Lock, Thread, Event
from datetime import datetime, timezone
import base64
import csv
import io
import json
import shutil
import time
import uuid

try:
    from openpyxl import load_workbook, Workbook
except Exception:  # pragma: no cover
    load_workbook = None
    Workbook = None

from execution_engine import derive_variable_key


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def safe_name(value: str) -> str:
    value = "".join(c if c.isalnum() or c in "-_." else "_" for c in str(value or ""))
    return value[:160] or "dataset"


class DataManager:
    """Phase 6 workbook/CSV storage, mapping, and sequential row execution."""

    TERMINAL = {"completed", "failed", "cancelled"}

    def __init__(self, root: Path, write_json, executor, recording_provider):
        self.root = Path(root)
        self.write_json = write_json
        self.executor = executor
        self.recording_provider = recording_provider
        self.data_root = self.root / "data" / "datasets"
        self.mapping_file = self.root / "data" / "data_mappings.json"
        self.batch_root = self.root / "data" / "batches"
        self.data_root.mkdir(parents=True, exist_ok=True)
        self.batch_root.mkdir(parents=True, exist_ok=True)
        if not self.mapping_file.exists():
            self.write_json(self.mapping_file, [])
        self.lock = Lock()
        self.thread = None
        self.cancel_event = Event()
        self.active_batch = None

    def dependency_status(self):
        return load_workbook is not None and Workbook is not None

    def _dataset_meta_path(self, dataset_id):
        return self.data_root / dataset_id / "dataset.json"

    def _batch_path(self, batch_id):
        return self.batch_root / batch_id / "batch.json"

    def _read_json(self, path, fallback):
        try:
            return json.loads(Path(path).read_text(encoding="utf-8"))
        except Exception:
            return fallback

    def list_datasets(self):
        rows = []
        for p in sorted(self.data_root.glob("*/dataset.json"), key=lambda x: x.stat().st_mtime, reverse=True):
            d = self._read_json(p, None)
            if d:
                rows.append(d)
        return rows

    def get_dataset(self, dataset_id, preview_rows=30, sheet=None):
        meta = self._read_json(self._dataset_meta_path(dataset_id), None)
        if not meta:
            return None
        try:
            preview = self.read_table(meta, sheet=sheet or meta.get("default_sheet"), limit=preview_rows)
        except Exception as exc:
            preview = {"columns": [], "rows": [], "error": str(exc)}
        return {**meta, "preview": preview}

    def upload(self, name: str, content_b64: str):
        if not name or not content_b64:
            raise ValueError("File name and content are required.")
        ext = Path(name).suffix.lower()
        if ext not in {".xlsx", ".csv"}:
            raise ValueError("Only .xlsx and .csv files are supported in Phase 6.")
        raw = base64.b64decode(content_b64)
        if len(raw) > 50 * 1024 * 1024:
            raise ValueError("Dataset exceeds the 50 MB Phase 6 limit.")
        dataset_id = str(uuid.uuid4())
        folder = self.data_root / dataset_id
        folder.mkdir(parents=True, exist_ok=True)
        source_name = "source" + ext
        source = folder / source_name
        source.write_bytes(raw)
        sheets = []
        default_sheet = None
        if ext == ".xlsx":
            if load_workbook is None:
                raise RuntimeError("openpyxl is not installed. Run pip install -r requirements.txt")
            wb = load_workbook(source, read_only=True, data_only=False)
            sheets = list(wb.sheetnames)
            default_sheet = sheets[0] if sheets else None
            wb.close()
        else:
            sheets = ["CSV"]
            default_sheet = "CSV"
        meta = {
            "schema_version": "webflow-dataset/1",
            "id": dataset_id,
            "name": name,
            "type": ext.lstrip("."),
            "source_path": "/" + source.relative_to(self.root).as_posix(),
            "source_file": source_name,
            "size_bytes": len(raw),
            "sheets": sheets,
            "default_sheet": default_sheet,
            "created_at": utc_now(),
            "updated_at": utc_now(),
        }
        self.write_json(self._dataset_meta_path(dataset_id), meta)
        return self.get_dataset(dataset_id)

    def _source_path(self, meta):
        return self.root / str(meta["source_path"]).lstrip("/")

    def read_table(self, meta, sheet=None, limit=None):
        source = self._source_path(meta)
        if meta.get("type") == "csv":
            with source.open("r", encoding="utf-8-sig", newline="") as fh:
                reader = csv.DictReader(fh)
                cols = list(reader.fieldnames or [])
                rows = []
                for i, row in enumerate(reader, start=2):
                    row["__row_number__"] = i
                    rows.append(row)
                    if limit and len(rows) >= limit:
                        break
                return {"sheet": "CSV", "columns": cols, "rows": rows}
        if load_workbook is None:
            raise RuntimeError("openpyxl is required for Excel files.")
        wb = load_workbook(source, read_only=True, data_only=False)
        ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        values = ws.iter_rows(values_only=True)
        try:
            headers_raw = next(values)
        except StopIteration:
            wb.close(); return {"sheet": ws.title, "columns": [], "rows": []}
        cols = []
        seen = {}
        for idx, value in enumerate(headers_raw, 1):
            base = str(value).strip() if value not in (None, "") else f"Column {idx}"
            count = seen.get(base, 0) + 1; seen[base] = count
            cols.append(base if count == 1 else f"{base} ({count})")
        rows = []
        for excel_row, values_row in enumerate(values, start=2):
            row = {cols[i]: values_row[i] if i < len(values_row) else None for i in range(len(cols))}
            row["__row_number__"] = excel_row
            rows.append(row)
            if limit and len(rows) >= limit:
                break
        title = ws.title
        wb.close()
        return {"sheet": title, "columns": cols, "rows": rows}

    def infer_variables(self, recording):
        result = []
        counters = {"Input": 0, "Secret": 0, "Output": 0, "Assertion": 0}
        for step in recording.get("steps", []):
            cls = step.get("classification")
            if cls not in counters:
                continue
            counters[cls] += 1
            target = step.get("target") or {}
            label = step.get("training_title") or target.get("label") or target.get("aria_label") or target.get("text") or target.get("name") or target.get("tag") or f"{cls} {counters[cls]}"
            key = derive_variable_key(step, counters[cls])
            result.append({
                "step_id": step.get("id"), "classification": cls, "key": key,
                "label": str(label)[:160], "action": step.get("action"),
                "required": cls in {"Secret"},
                "recorded_value": None if cls == "Secret" else step.get("value"),
            })
        return result

    def list_mappings(self, recording_id=None):
        rows = self._read_json(self.mapping_file, [])
        if recording_id:
            rows = [m for m in rows if m.get("recording_id") == recording_id]
        return rows

    def save_mapping(self, payload: dict, recording: dict):
        dataset_id = str(payload.get("dataset_id") or "")
        dataset = self.get_dataset(dataset_id, preview_rows=1)
        if not dataset:
            raise ValueError("Dataset not found.")
        mapping_id = str(payload.get("id") or uuid.uuid4())
        variables = self.infer_variables(recording)
        valid_keys = {v["key"] for v in variables}
        input_map = {str(k): str(v) for k, v in (payload.get("input_map") or {}).items() if k in valid_keys and str(v)}
        output_map = {str(k): str(v) for k, v in (payload.get("output_map") or {}).items() if k in valid_keys and str(v)}
        mapping = {
            "schema_version": "webflow-data-mapping/1",
            "id": mapping_id,
            "name": str(payload.get("name") or "Data Mapping"),
            "recording_id": recording.get("id"),
            "flow_id": recording.get("flow_id"),
            "dataset_id": dataset_id,
            "sheet": str(payload.get("sheet") or dataset.get("default_sheet") or ""),
            "input_map": input_map,
            "output_map": output_map,
            "result_column": str(payload.get("result_column") or "WebFlow Result"),
            "error_column": str(payload.get("error_column") or "WebFlow Error"),
            "run_id_column": str(payload.get("run_id_column") or "WebFlow Run ID"),
            "created_at": utc_now(), "updated_at": utc_now(),
        }
        rows = self._read_json(self.mapping_file, [])
        existing = next((i for i, m in enumerate(rows) if m.get("id") == mapping_id), None)
        if existing is None:
            rows.insert(0, mapping)
        else:
            mapping["created_at"] = rows[existing].get("created_at", mapping["created_at"])
            rows[existing] = mapping
        self.write_json(self.mapping_file, rows)
        return mapping

    def status(self):
        with self.lock:
            return json.loads(json.dumps(self.active_batch)) if self.active_batch else {"state": "idle"}

    def list_batches(self, limit=30):
        items = []
        for p in sorted(self.batch_root.glob("*/batch.json"), key=lambda x: x.stat().st_mtime, reverse=True)[:limit]:
            b = self._read_json(p, None)
            if not b: continue
            items.append({k: b.get(k) for k in ["id","state","recording_id","flow_id","dataset_id","mapping_id","started_at","ended_at","total_rows","completed_rows","failed_rows","result_path","parent_batch_id"]})
        return items

    def get_batch(self, batch_id):
        if self.active_batch and self.active_batch.get("id") == batch_id:
            return self.status()
        return self._read_json(self._batch_path(batch_id), None)

    def _persist_batch(self):
        if self.active_batch:
            self.write_json(self._batch_path(self.active_batch["id"]), self.active_batch)

    def start_batch(self, recording, mapping, dataset, *, row_numbers=None, headless=True, timeout_ms=12000, retries=1, failure_policy="continue", parent_batch_id=None):
        if not self.dependency_status():
            return False, {"error": "openpyxl_not_installed", "message": "Install Phase 6 dependencies first."}
        with self.lock:
            if self.thread and self.thread.is_alive():
                return False, {"error": "batch_busy", "message": "A data-driven batch is already active."}
            batch_id = str(uuid.uuid4())
            self.active_batch = {
                "schema_version": "webflow-batch/1", "id": batch_id,
                "recording_id": recording.get("id"), "flow_id": recording.get("flow_id"),
                "dataset_id": dataset.get("id"), "mapping_id": mapping.get("id"),
                "sheet": mapping.get("sheet"), "state": "starting", "started_at": utc_now(), "ended_at": None,
                "headless": bool(headless), "timeout_ms": int(timeout_ms or 12000), "retries": int(retries or 0),
                "failure_policy": failure_policy, "requested_rows": list(row_numbers or []),
                "total_rows": 0, "completed_rows": 0, "failed_rows": 0, "current_row": None,
                "rows": [], "result_path": None, "last_error": None, "parent_batch_id": parent_batch_id,
            }
            self.cancel_event.clear(); self._persist_batch()
            self.thread = Thread(target=self._run_batch, args=(recording, mapping, dataset), daemon=True, name="WebFlowDataBatch")
            self.thread.start()
            return True, self.status()

    def cancel_batch(self):
        with self.lock:
            if not (self.thread and self.thread.is_alive()):
                return False, {"error": "no_active_batch"}
            self.active_batch["state"] = "cancelling"; self._persist_batch()
        self.cancel_event.set()
        try:
            self.executor.cancel()
        except Exception:
            pass
        return True, self.status()

    def retry_failed(self, batch_id, recording, mapping, dataset):
        old = self.get_batch(batch_id)
        if not old:
            return False, {"error": "batch_not_found"}
        failed = [r.get("row_number") for r in old.get("rows", []) if r.get("status") == "failed"]
        if not failed:
            return False, {"error": "no_failed_rows", "message": "This batch has no failed rows."}
        return self.start_batch(recording, mapping, dataset, row_numbers=failed, headless=old.get("headless", True), timeout_ms=old.get("timeout_ms",12000), retries=old.get("retries",1), failure_policy=old.get("failure_policy","continue"), parent_batch_id=batch_id)

    def _run_batch(self, recording, mapping, dataset):
        try:
            table = self.read_table(dataset, sheet=mapping.get("sheet"), limit=None)
            rows = table["rows"]
            requested = {int(x) for x in self.active_batch.get("requested_rows", []) if str(x).isdigit()}
            if requested:
                rows = [r for r in rows if int(r.get("__row_number__", 0)) in requested]
            with self.lock:
                self.active_batch["total_rows"] = len(rows); self.active_batch["state"] = "running"; self._persist_batch()
            result_rows = {}
            for row in rows:
                if self.cancel_event.is_set():
                    with self.lock: self.active_batch["state"] = "cancelled"; self._persist_batch()
                    break
                row_number = int(row.get("__row_number__", 0))
                variables = {}
                for key, col in mapping.get("input_map", {}).items():
                    variables[key] = row.get(col)
                row_result = {"row_number": row_number, "status": "running", "run_id": None, "outputs": {}, "error": None, "started_at": utc_now(), "ended_at": None}
                with self.lock:
                    self.active_batch["current_row"] = row_number; self.active_batch["rows"].append(row_result); self._persist_batch()
                ok, run = self.executor.start(recording, flow_id=recording.get("flow_id"), headless=self.active_batch.get("headless", True), timeout_ms=self.active_batch.get("timeout_ms",12000), retries=self.active_batch.get("retries",1), failure_policy="stop", variables=variables, batch_context={"batch_id":self.active_batch["id"],"row_number":row_number})
                if not ok:
                    row_result["status"] = "failed"; row_result["error"] = run.get("message") or run.get("error"); row_result["ended_at"] = utc_now()
                else:
                    row_result["run_id"] = run.get("id")
                    while True:
                        if self.cancel_event.is_set():
                            self.executor.cancel()
                        status = self.executor.status()
                        if status.get("id") == run.get("id") and status.get("state") in self.TERMINAL:
                            run = status
                            # The status becomes terminal just before the worker thread exits.
                            # Wait briefly so the next workbook row cannot collide with the prior worker.
                            for _ in range(40):
                                t = getattr(self.executor, "thread", None)
                                if not (t and t.is_alive()): break
                                time.sleep(.05)
                            break
                        time.sleep(.25)
                    row_result["status"] = "completed" if run.get("state") == "completed" else "failed"
                    row_result["outputs"] = run.get("outputs") or {}
                    row_result["error"] = run.get("last_error")
                    row_result["ended_at"] = utc_now()
                result_rows[row_number] = row_result
                with self.lock:
                    if row_result["status"] == "completed": self.active_batch["completed_rows"] += 1
                    else: self.active_batch["failed_rows"] += 1
                    self._persist_batch()
                if row_result["status"] == "failed" and self.active_batch.get("failure_policy") == "stop":
                    with self.lock: self.active_batch["state"] = "failed"; self._persist_batch()
                    break
            result_path = self._write_results(dataset, mapping, result_rows, self.active_batch["id"])
            with self.lock:
                self.active_batch["result_path"] = result_path
                self.active_batch["current_row"] = None
                if self.active_batch.get("state") in {"running","starting","cancelling"}:
                    self.active_batch["state"] = "failed" if self.active_batch.get("failed_rows") else "completed"
        except Exception as exc:
            with self.lock:
                self.active_batch["state"] = "failed"; self.active_batch["last_error"] = f"{type(exc).__name__}: {exc}"
        finally:
            with self.lock:
                self.active_batch["ended_at"] = utc_now(); self._persist_batch()

    def _write_results(self, dataset, mapping, result_rows, batch_id):
        if Workbook is None or load_workbook is None:
            return None
        folder = self.batch_root / batch_id
        folder.mkdir(parents=True, exist_ok=True)
        out = folder / "results.xlsx"
        source = self._source_path(dataset)
        if dataset.get("type") == "xlsx":
            shutil.copy2(source, out)
            wb = load_workbook(out)
            ws = wb[mapping.get("sheet")] if mapping.get("sheet") in wb.sheetnames else wb[wb.sheetnames[0]]
            headers = [str(c.value).strip() if c.value not in (None,"") else f"Column {i}" for i,c in enumerate(ws[1],1)]
        else:
            table = self.read_table(dataset, sheet="CSV", limit=None)
            wb = Workbook(); ws = wb.active; ws.title = "Results"; headers = list(table["columns"])
            ws.append(headers)
            for r in table["rows"]:
                ws.append([r.get(c) for c in headers])
        def ensure_col(name):
            if name in headers: return headers.index(name)+1
            headers.append(name); col=len(headers); ws.cell(1,col,name); return col
        result_col=ensure_col(mapping.get("result_column") or "WebFlow Result")
        error_col=ensure_col(mapping.get("error_column") or "WebFlow Error")
        run_col=ensure_col(mapping.get("run_id_column") or "WebFlow Run ID")
        out_cols={key:ensure_col(col) for key,col in mapping.get("output_map",{}).items()}
        for row_number, rr in result_rows.items():
            ws.cell(row_number,result_col,rr.get("status"))
            ws.cell(row_number,error_col,rr.get("error") or "")
            ws.cell(row_number,run_col,rr.get("run_id") or "")
            for key,col_idx in out_cols.items():
                value=(rr.get("outputs") or {}).get(key)
                ws.cell(row_number,col_idx,value)
        wb.save(out); wb.close()
        return "/" + out.relative_to(self.root).as_posix()
